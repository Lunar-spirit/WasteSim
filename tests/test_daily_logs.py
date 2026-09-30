import io

BOUNDARY = {
    "type": "MultiPolygon",
    "coordinates": [[[[74.79, 13.34], [74.81, 13.34], [74.81, 13.36], [74.79, 13.36], [74.79, 13.34]]]],
}


async def _create_habitation(client, headers, name="LogVille") -> str:
    resp = await client.post(
        "/api/v1/habitations",
        headers=headers,
        json={
            "name": name,
            "habitation_type": "VILLAGE",
            "state": "Karnataka",
            "district": "Udupi",
            "boundary_geojson": BOUNDARY,
        },
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["data"]["id"]


VALID_LOG = {
    "log_date": "2026-03-01",
    "total_collected_tonnes": 12.5,
    "organic_tonnes": 7.0,
    "dry_recyclable_tonnes": 4.5,
    "hazardous_tonnes": 0.2,
    "vehicles_deployed": 3,
    "trips_completed": 6,
    "diesel_consumed_litres": 40.0,
    "collection_coverage_pct_observed": 82.5,
    "anomaly_flag": "NORMAL",
    "notes": "Routine day",
}


async def test_create_then_update_same_day_upserts(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "UpsertVille")

    create_resp = await client.post(
        f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=VALID_LOG
    )
    assert create_resp.status_code == 200, create_resp.text
    created = create_resp.json()["data"]
    assert created["created"] is True
    assert created["total_collected_tonnes"] == 12.5

    updated_payload = {**VALID_LOG, "total_collected_tonnes": 15.0, "anomaly_flag": "FESTIVAL_SURGE"}
    update_resp = await client.post(
        f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=updated_payload
    )
    assert update_resp.status_code == 200, update_resp.text
    updated = update_resp.json()["data"]
    assert updated["created"] is False
    assert updated["total_collected_tonnes"] == 15.0
    assert updated["anomaly_flag"] == "FESTIVAL_SURGE"
    assert updated["id"] == created["id"]

    list_resp = await client.get(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers)
    assert list_resp.json()["data"]["total"] == 1


async def test_create_rejects_negative_values(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "NegativeVille")
    bad_payload = {**VALID_LOG, "organic_tonnes": -5.0}
    resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=bad_payload)
    # This project maps RequestValidationError to 400, not FastAPI's default
    # 422 (app/core/errors.py) — matches every other malformed-body test.
    assert resp.status_code == 400, resp.text


async def test_unassigned_planner_cannot_log(client, admin_headers, planner_headers):
    habitation_id = await _create_habitation(client, admin_headers, "UnassignedPlannerVille")
    resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=planner_headers, json=VALID_LOG)
    assert resp.status_code == 403, resp.text
    assert resp.json()["error"]["code"] == "FORBIDDEN_RESOURCE"


async def test_assigned_planner_can_log(client, admin_headers, planner_headers, planner_user):
    habitation_id = await _create_habitation(client, admin_headers, "AssignedPlannerVille")
    planner, _ = planner_user
    assign_resp = await client.post(
        "/api/v1/admin/access/assign",
        headers=admin_headers,
        json={"user_id": str(planner.id), "habitation_id": habitation_id},
    )
    assert assign_resp.status_code == 201, assign_resp.text

    resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=planner_headers, json=VALID_LOG)
    assert resp.status_code == 200, resp.text
    assert resp.json()["data"]["logged_by"] == str(planner.id)


async def test_researcher_can_read_but_not_write(client, admin_headers, researcher_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ResearcherLogVille")
    await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=VALID_LOG)

    read_resp = await client.get(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=researcher_headers)
    assert read_resp.status_code == 200, read_resp.text
    assert read_resp.json()["data"]["total"] == 1

    write_resp = await client.post(
        f"/api/v1/habitations/{habitation_id}/daily-logs", headers=researcher_headers, json=VALID_LOG
    )
    assert write_resp.status_code == 403, write_resp.text


async def test_list_filters_by_date_range_and_paginates(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "DateRangeVille")
    for day in (1, 2, 3, 4, 5):
        payload = {**VALID_LOG, "log_date": f"2026-03-0{day}"}
        resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=payload)
        assert resp.status_code == 200, resp.text

    ranged = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs",
        headers=admin_headers,
        params={"from_date": "2026-03-02", "to_date": "2026-03-04"},
    )
    assert ranged.status_code == 200, ranged.text
    data = ranged.json()["data"]
    assert data["total"] == 3
    assert {row["log_date"] for row in data["items"]} == {"2026-03-02", "2026-03-03", "2026-03-04"}

    paged = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs",
        headers=admin_headers,
        params={"page": 1, "page_size": 2},
    )
    paged_data = paged.json()["data"]
    assert paged_data["total"] == 5
    assert len(paged_data["items"]) == 2
    # newest first
    assert paged_data["items"][0]["log_date"] == "2026-03-05"


CSV_VALID = """log_date,organic_tonnes,dry_recyclable_tonnes,vehicles_deployed,trips_completed,hazardous_tonnes,anomaly_flag
2026-04-01,6.0,3.0,2,4,0.1,NORMAL
2026-04-02,6.5,3.2,2,4,,MONSOON_FLOOD
"""

CSV_WITH_BAD_ROW = """log_date,organic_tonnes,dry_recyclable_tonnes,vehicles_deployed,trips_completed
2026-04-10,5.0,2.0,2,4
2026-04-11,not-a-number,2.0,2,4
2026-04-12,4.0,1.0,-1,4
"""

CSV_MISSING_COLUMN = """organic_tonnes,dry_recyclable_tonnes,vehicles_deployed,trips_completed
5.0,2.0,2,4
"""


async def test_bulk_csv_derives_total_and_imports_valid_rows(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "BulkCsvVille")
    files = {"file": ("logs.csv", io.BytesIO(CSV_VALID.encode()), "text/csv")}
    resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs/bulk-csv", headers=admin_headers, files=files)
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert data["total_rows"] == 2
    assert data["created_count"] == 2
    assert data["error_count"] == 0

    list_resp = await client.get(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers)
    rows = {row["log_date"]: row for row in list_resp.json()["data"]["items"]}
    # Automatic summary derivation: total_collected_tonnes wasn't in the CSV at all.
    assert rows["2026-04-01"]["total_collected_tonnes"] == 9.1
    assert rows["2026-04-02"]["total_collected_tonnes"] == 9.7
    assert rows["2026-04-02"]["anomaly_flag"] == "MONSOON_FLOOD"


async def test_bulk_csv_reports_per_row_errors_without_aborting_the_batch(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "PartialFailVille")
    files = {"file": ("logs.csv", io.BytesIO(CSV_WITH_BAD_ROW.encode()), "text/csv")}
    resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs/bulk-csv", headers=admin_headers, files=files)
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert data["total_rows"] == 3
    assert data["created_count"] == 1
    assert data["error_count"] == 2
    assert {e["row_number"] for e in data["errors"]} == {2, 3}


async def test_bulk_csv_missing_required_column_rejected(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "MissingColumnVille")
    files = {"file": ("logs.csv", io.BytesIO(CSV_MISSING_COLUMN.encode()), "text/csv")}
    resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs/bulk-csv", headers=admin_headers, files=files)
    assert resp.status_code == 422, resp.text
    assert resp.json()["error"]["code"] == "DAILY_LOG_CSV_MISSING_COLUMNS"


async def test_export_requires_admin(client, admin_headers, planner_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ExportAuthVille")
    resp = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs/export",
        headers=planner_headers,
        params={"period_type": "yearly", "year": 2026},
    )
    assert resp.status_code == 403, resp.text
    assert resp.json()["error"]["code"] == "FORBIDDEN_ROLE"


async def test_export_yearly_requires_year(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ExportNoYearVille")
    resp = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs/export",
        headers=admin_headers,
        params={"period_type": "yearly"},
    )
    assert resp.status_code == 400, resp.text
    assert resp.json()["error"]["code"] == "EXPORT_YEAR_REQUIRED"


async def test_export_monthly_requires_year_and_month(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ExportNoMonthVille")
    resp = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs/export",
        headers=admin_headers,
        params={"period_type": "monthly", "year": 2026},
    )
    assert resp.status_code == 400, resp.text
    assert resp.json()["error"]["code"] == "EXPORT_YEAR_MONTH_REQUIRED"


async def test_export_custom_requires_both_dates(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ExportNoRangeVille")
    resp = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs/export",
        headers=admin_headers,
        params={"period_type": "custom", "start_date": "2026-01-01"},
    )
    assert resp.status_code == 400, resp.text
    assert resp.json()["error"]["code"] == "EXPORT_RANGE_REQUIRED"


async def test_export_monthly_csv_filters_correctly_and_includes_summary_row(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ExportMonthlyVille")
    # Two days inside March, one outside (April) — must not appear in a March export.
    for log_date, total in (("2026-03-05", 10.0), ("2026-03-15", 14.0), ("2026-04-01", 99.0)):
        payload = {**VALID_LOG, "log_date": log_date, "total_collected_tonnes": total}
        resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=payload)
        assert resp.status_code == 200, resp.text

    resp = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs/export",
        headers=admin_headers,
        params={"period_type": "monthly", "year": 2026, "month": 3},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    assert "ExportMonthlyVille" in resp.headers["content-disposition"]
    assert "2026_03" in resp.headers["content-disposition"]

    lines = resp.text.strip().splitlines()
    header = lines[0].split(",")
    assert header == [
        "Date",
        "Habitation Name",
        "Logged By (Email)",
        "Total Collected (t)",
        "Organic (t)",
        "Dry Recyclable (t)",
        "Hazardous (t)",
        "Vehicles Deployed",
        "Trips Completed",
        "Diesel Consumed (L)",
        "Observed Coverage (%)",
        "Anomaly / Notes",
    ]
    # header + 2 March rows + 1 summary row, April excluded entirely
    assert len(lines) == 4
    assert "2026-03-05" in lines[1]
    assert "2026-03-15" in lines[2]
    assert "99.0" not in resp.text  # the April row's total never appears

    summary_row = lines[3].split(",")
    assert summary_row[0] == "TOTAL / AVERAGE"
    assert float(summary_row[3]) == 24.0  # 10.0 + 14.0


async def test_export_custom_range_and_xlsx_format(client, admin_headers):
    habitation_id = await _create_habitation(client, admin_headers, "ExportXlsxVille")
    for log_date in ("2026-06-01", "2026-06-10"):
        payload = {**VALID_LOG, "log_date": log_date}
        resp = await client.post(f"/api/v1/habitations/{habitation_id}/daily-logs", headers=admin_headers, json=payload)
        assert resp.status_code == 200, resp.text

    resp = await client.get(
        f"/api/v1/habitations/{habitation_id}/daily-logs/export",
        headers=admin_headers,
        params={"period_type": "custom", "start_date": "2026-06-01", "end_date": "2026-06-30", "format": "xlsx"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.content[:2] == b"PK"  # xlsx is a real zip archive, not a renamed CSV

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Date"
    assert len(rows) == 4  # header + 2 logged days + summary row
    assert rows[-1][0] == "TOTAL / AVERAGE"
